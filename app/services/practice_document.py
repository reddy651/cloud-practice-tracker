from datetime import datetime
from io import BytesIO


def _ordered(practice):
    return sorted(practice.screenshots, key=lambda x: ((x.sequence_no or 999999), (x.image_order or 999999), x.id))

def _groups(practice):
    groups = {}
    for s in _ordered(practice):
        groups.setdefault(s.sequence_no or 0, []).append(s)
    return [(seq, rows) for seq, rows in sorted(groups.items())]


def build_markdown(practice):
    lines = [
        f"# {practice.title}", "",
        f"- **Provider:** {practice.provider}",
        f"- **Service:** {practice.service or '-'}",
        f"- **Topic:** {practice.topic or '-'}", "",
        "## Practice Steps", "",
    ]
    for seq, rows in _groups(practice):
        explanation = next((r.step_explanation for r in rows if (r.step_explanation or '').strip()), 'Not provided yet.')
        lines += [f"### Step #{seq:03d}", "**Step explanation:**", explanation, ""]
        for i, row in enumerate(rows, start=1):
            lines += [f"- Screenshot {i:03d}: `{row.original_name or 'Replacement required'}`"]
        lines.append("")
    lines += ["## Additional Notes", "", practice.description or "", "", f"Generated/updated: {datetime.now()}"]
    return "\n".join(lines)


def build_docx(practice, screenshot_items=None):
    """Build a Word document grouped by step; each step may contain multiple screenshots."""
    from docx import Document
    from docx.shared import Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    title = doc.add_heading("Cloud Practice Tracker", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading(practice.title, level=1)
    p = doc.add_paragraph()
    p.add_run("Provider: ").bold = True; p.add_run(practice.provider or "-")
    p.add_run("\nService: ").bold = True; p.add_run(practice.service or "-")
    p.add_run("\nTopic: ").bold = True; p.add_run(practice.topic or "-")

    items = screenshot_items or []
    item_by_id = {item.get("id"): item for item in items if item.get("id") is not None}
    groups = _groups(practice)
    doc.add_paragraph(f"Total steps: {len(groups)}")
    doc.add_heading("Practice Steps", level=2)

    for seq, rows in groups:
        doc.add_heading(f"Step #{seq:03d}", level=3)
        ep = doc.add_paragraph()
        ep.add_run(rows[0].step_explanation or "Not provided yet.")
        for row in rows:
            item = item_by_id.get(row.id)
            if item and item.get("data"):
                try:
                    img_para = doc.add_paragraph(); img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    img_para.add_run().add_picture(BytesIO(item["data"]), width=Inches(6.4))
                except Exception:
                    doc.add_paragraph("[Screenshot could not be embedded]")
            else:
                doc.add_paragraph("[Screenshot replacement required]")
        doc.add_paragraph()

    if practice.description:
        doc.add_heading("Additional Notes", level=2)
        doc.add_paragraph(practice.description)
    return _docx_bytes(doc)

def _docx_bytes(document):
    out = BytesIO()
    document.save(out)
    return out.getvalue()
